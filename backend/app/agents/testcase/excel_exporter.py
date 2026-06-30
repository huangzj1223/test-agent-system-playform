"""
测试用例 Excel 导出工具

本模块提供将测试用例导出为 Excel 文件的能力，支持企业级测试管理工具
（如禅道、Tapd、TestRail）的导入格式。
"""

from pathlib import Path
from typing import Any

from openpyxl import Workbook

# 与 tools.EXPORTS_DIR 保持一致：项目根目录 exports/
_EXPORTS_DIR = Path(__file__).resolve().parents[4] / "exports"
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# Excel 样式配置
_HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_ALIGNMENT_WRAP = Alignment(vertical="top", wrap_text=True)
_ALIGNMENT_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 默认列宽配置
_DEFAULT_COLUMN_WIDTHS = {
    "A": 18,  # 用例编号
    "B": 35,  # 用例标题
    "C": 14,  # 所属模块
    "D": 12,  # 用例类型
    "E": 10,  # 优先级
    "F": 30,  # 前置条件
    "G": 40,  # 测试步骤
    "H": 30,  # 测试数据
    "I": 40,  # 预期结果
    "J": 20,  # 备注
}


def _flatten_steps(steps: list[dict[str, Any]] | str | None) -> str:
    """将步骤列表转换为带序号的文本。"""
    if not steps:
        return ""
    # 如果已经是字符串，直接返回
    if isinstance(steps, str):
        return steps
    # 如果是列表，进行格式化
    lines = []
    for step in steps:
        # 防止step是字符串而不是字典的情况
        if isinstance(step, str):
            lines.append(step)
            continue
        seq = step.get("seq", step.get("step", len(lines) + 1))
        action = step.get("action", step.get("操作描述", ""))
        target = step.get("target", step.get("操作对象", ""))
        data = step.get("data", "")
        line = f"{seq}. {action}"
        if target:
            line += f" [{target}]"
        if data:
            line += f"（数据：{data}）"
        lines.append(line)
    return "\n".join(lines)


def _flatten_test_data(test_data: Any) -> str:
    """将测试数据转换为文本。"""
    if not test_data:
        return ""
    if isinstance(test_data, str):
        return test_data
    if isinstance(test_data, dict):
        lines = [f"{k}: {v}" for k, v in test_data.items()]
        return "\n".join(lines)
    if isinstance(test_data, list):
        lines = []
        for idx, item in enumerate(test_data, start=1):
            if isinstance(item, dict):
                lines.append("; ".join(f"{k}: {v}" for k, v in item.items()))
            else:
                lines.append(f"{idx}. {item}")
        return "\n".join(lines)
    return str(test_data)


def _flatten_expected_results(expected_results: list[str] | str | None) -> str:
    """将预期结果列表转换为文本。"""
    if not expected_results:
        return ""
    if isinstance(expected_results, str):
        return expected_results
    # 防止列表中包含非字符串元素
    lines = []
    for idx, result in enumerate(expected_results, start=1):
        if isinstance(result, str):
            lines.append(f"{idx}. {result}")
        else:
            # 如果不是字符串，转换为字符串
            lines.append(f"{idx}. {str(result)}")
    return "\n".join(lines)


def _flatten_preconditions(preconditions: list[str] | str | None) -> str:
    """将前置条件列表转换为文本。"""
    if not preconditions:
        return ""
    if isinstance(preconditions, str):
        return preconditions
    # 防止列表中包含非字符串元素
    lines = []
    for idx, cond in enumerate(preconditions, start=1):
        if isinstance(cond, str):
            lines.append(f"{idx}. {cond}")
        else:
            # 如果不是字符串，转换为字符串
            lines.append(f"{idx}. {str(cond)}")
    return "\n".join(lines)


def _extract_field(case: dict[str, Any], *keys: str, default: Any = "") -> Any:
    """从字典中按多个候选键提取值。"""
    for key in keys:
        if key in case:
            return case[key]
    return default


def _first_non_empty(case: dict[str, Any], *keys: str, default: Any = "") -> Any:
    """按候选键提取第一个非空值，避免空别名遮蔽有效字段。"""
    for key in keys:
        if key not in case:
            continue
        value = case[key]
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        if isinstance(value, (list, dict)) and not value:
            continue
        return value
    return default


_EXCEL_HEADERS = [
    "用例编号",
    "用例标题",
    "所属模块",
    "用例类型",
    "优先级",
    "前置条件",
    "测试步骤",
    "测试数据",
    "预期结果",
    "备注",
]


def resolve_excel_output_path(output_path: str | Path, *, prefer_existing: bool = False) -> Path:
    """
    将导出路径规范为绝对路径，避免工作目录不同导致「追加写到另一个文件」。

    - 绝对路径：原样 resolve
    - 相对路径 exports/xxx.xlsx 或 xxx.xlsx：统一落到项目根 exports/ 目录
    - prefer_existing=True 时，若目标不存在则按文件名在 exports/ 中查找已有文件（便于 append）
    """
    raw = Path(output_path)
    if raw.is_absolute():
        resolved = raw.resolve()
    else:
        parts = raw.parts
        if parts and parts[0].lower() == "exports":
            resolved = (_EXPORTS_DIR / Path(*parts[1:])).resolve()
        else:
            resolved = (_EXPORTS_DIR / raw).resolve()

    if prefer_existing and not resolved.exists():
        by_name = _EXPORTS_DIR / raw.name
        if by_name.exists():
            return by_name.resolve()
        matches = sorted(_EXPORTS_DIR.glob(raw.name))
        if len(matches) == 1:
            return matches[0].resolve()
    return resolved


def _save_workbook(wb, output_path: Path) -> Path:
    """保存工作簿；文件被 Excel 占用时明确报错，避免静默写到带时间戳的新文件。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(str(output_path))
    except PermissionError as exc:
        raise PermissionError(
            f"无法写入 {output_path}，请先关闭 Excel 中打开的该文件后再重试导出。"
        ) from exc
    return output_path.resolve()


def _sanitize_sheet_name(name: str) -> str:
    """Excel 工作表名称：最长 31 字符，且不能含 : \\ / ? * [ ]"""
    cleaned = (name or "").strip()
    for ch in (":", "\\", "/", "?", "*", "[", "]"):
        cleaned = cleaned.replace(ch, "_")
    return (cleaned[:31] or "测试用例")


def _format_case_id(raw_id: Any, index: int) -> str:
    """返回用例编号；缺失时生成稳定的默认编号。"""
    case_id = str(raw_id).strip() if raw_id is not None else ""
    return case_id or f"TC-{index:03d}"


def _write_case_rows(ws, test_cases: list[dict[str, Any]], start_index: int = 1) -> None:
    """向工作表追加用例行（不含表头）。"""
    headers_count = 10
    for offset, case in enumerate(test_cases):
        index = start_index + offset
        row = [
            _format_case_id(_first_non_empty(case, "id", "case_id", "用例编号", "用例标识"), index),
            _first_non_empty(case, "title", "name", "case_name", "用例标题", "用例名称", "测试用例名称"),
            _first_non_empty(case, "module", "所属模块", "功能模块", "模块", "追踪关系"),
            _first_non_empty(case, "type", "case_type", "用例类型", "测试类型", "设计方法"),
            _first_non_empty(case, "priority", "优先级", "级别"),
            _flatten_preconditions(
                _first_non_empty(
                    case,
                    "preconditions",
                    "precondition",
                    "前置条件",
                    "前提条件",
                    "前提与约束",
                    "前提和约束",
                    "用例初始化",
                    default=None,
                )
            ),
            _flatten_steps(_first_non_empty(case, "steps", "测试步骤", "操作步骤", "输入及操作", default=None)),
            _flatten_test_data(
                _first_non_empty(
                    case, "test_data", "testData", "测试数据", "具体测试数据", "测试输入", "输入数据", default=None
                )
            ),
            _flatten_expected_results(
                _first_non_empty(
                    case,
                    "expected_results",
                    "expected",
                    "预期结果",
                    "期望结果",
                    "期望结果与评估标准",
                    default=None,
                )
            ),
            _first_non_empty(case, "remarks", "备注", "关联需求", "用例说明", "测试用例综述"),
        ]
        ws.append(row)
        row_idx = ws.max_row
        for col_idx in range(1, headers_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = _ALIGNMENT_WRAP
            cell.border = _BORDER
        ws.row_dimensions[row_idx].height = 60


def _write_sheet_headers(ws) -> None:
    """写入表头行（仅当工作表为空时）。"""
    if ws.max_row >= 1 and ws.cell(row=1, column=1).value:
        return
    # 新建工作簿的 active 表第 1 行常为空，用 append 会把表头写到第 2 行
    if ws.max_row <= 1 and not ws.cell(row=1, column=1).value:
        for col_idx, header in enumerate(_EXCEL_HEADERS, start=1):
            ws.cell(row=1, column=col_idx, value=header)
        return
    ws.append(list(_EXCEL_HEADERS))


def _apply_header_style(ws) -> None:
    for col_idx, _header in enumerate(_EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _ALIGNMENT_CENTER
        cell.border = _BORDER
    ws.row_dimensions[1].height = 24
    for col_letter, width in _DEFAULT_COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width


def _get_or_create_sheet(wb, sheet_name: str):
    """按菜单名获取工作表；不存在则新建并写入表头。"""
    safe_name = _sanitize_sheet_name(sheet_name)
    if safe_name in wb.sheetnames:
        return wb[safe_name]
    ws = wb.create_sheet(title=safe_name)
    _write_sheet_headers(ws)
    _apply_header_style(ws)
    return ws


def _count_existing_cases(ws) -> int:
    """表头占第 1 行，返回已有用例行数（按用例编号列非空统计）。"""
    if ws.max_row <= 1:
        return 0
    count = 0
    for row_idx in range(2, ws.max_row + 1):
        cell_val = ws.cell(row=row_idx, column=1).value
        if cell_val is not None and str(cell_val).strip():
            count += 1
    return count


def count_cases_in_sheet(output_path: str | Path, sheet_name: str) -> int:
    """读取指定页签已有用例行数（不含表头）；文件或页签不存在时返回 0。"""
    path = resolve_excel_output_path(output_path, prefer_existing=True)
    if not path.exists():
        return 0
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    safe_name = _sanitize_sheet_name(sheet_name)
    if safe_name not in wb.sheetnames:
        return 0
    return _count_existing_cases(wb[safe_name])


def append_test_cases_to_excel(
    test_cases: list[dict[str, Any]],
    output_path: str | Path,
    sheet_name: str = "测试用例",
) -> str:
    """
    向已有 Excel 文件追加测试用例；文件不存在时自动创建（含表头）。

    Args:
        test_cases: 待追加的用例列表。
        output_path: Excel 文件路径（与首批导出使用同一路径）。
        sheet_name: 工作表名称。

    Returns:
        文件绝对路径。
    """
    if not test_cases:
        raise ValueError("测试用例列表为空，无法追加到 Excel。")

    output_path = resolve_excel_output_path(output_path, prefer_existing=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if not output_path.exists():
        return export_test_cases_to_excel(test_cases, output_path, sheet_name)

    from openpyxl import load_workbook

    wb = load_workbook(str(output_path))
    ws = _get_or_create_sheet(wb, sheet_name)
    before_count = _count_existing_cases(ws)
    start_index = before_count + 1
    _write_case_rows(ws, test_cases, start_index=start_index)
    after_count = _count_existing_cases(ws)
    added = after_count - before_count
    if added != len(test_cases):
        raise RuntimeError(
            f"追加写入异常：预期写入 {len(test_cases)} 条，实际增加 {added} 条。"
            f"请关闭 Excel 中打开的文件后重试，并确认 output_path 与首批导出路径一致。"
        )

    saved = _save_workbook(wb, output_path)
    return str(saved)


def export_test_cases_to_excel(
    test_cases: list[dict[str, Any]],
    output_path: str | Path,
    sheet_name: str = "测试用例",
) -> str:
    """
    将测试用例列表导出为 Excel 文件。

    支持的测试用例字段（兼容 JSON / CSV / Markdown 中定义的格式）：
      - id / 用例编号
      - title / 用例标题
      - module / 所属模块
      - type / 用例类型
      - priority / 优先级
      - preconditions / 前置条件
      - steps / 测试步骤
      - test_data / 测试数据
      - expected_results / 预期结果
      - remarks / 备注

    Args:
        test_cases: 测试用例字典列表，每个字典描述一条用例。
        output_path: 导出的 Excel 文件路径（支持 str 或 Path）。
        sheet_name: 工作表名称，默认为 "测试用例"。

    Returns:
        导出文件的绝对路径字符串。
    """
    if not test_cases:
        raise ValueError("测试用例列表为空，无法导出 Excel。")

    output_path = resolve_excel_output_path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("无法创建工作表。")
    ws.title = _sanitize_sheet_name(sheet_name)

    _write_sheet_headers(ws)
    _apply_header_style(ws)
    _write_case_rows(ws, test_cases, start_index=1)

    saved = _save_workbook(wb, output_path)
    return str(saved)
