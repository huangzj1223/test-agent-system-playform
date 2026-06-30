"""
测试用例 DOCX (Word) 导出工具

本模块提供将测试用例导出为 Word (.docx) 的功能。
为每个用例生成一个独立的表格，并设置指定的列宽比例和格式。
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml


def _set_cell_bg(cell, color="D9D9D9"):
    """设置单元格背景颜色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    tcPr.append(shd)


def _flatten_steps(steps: list[dict[str, Any]] | str | None) -> str:
    """将步骤列表转换为带序号的文本。"""
    if not steps:
        return ""
    # 如果已经是字符串，直接返回
    if isinstance(steps, str):
        return steps
    # 如果是列表，进行格式化
    lines = []
    for step in steps:
        # 防止step是字符串而不是字典的情况
        if isinstance(step, str):
            lines.append(step)
            continue
        seq = step.get("seq", step.get("step", len(lines) + 1))
        action = step.get("action", step.get("操作描述", ""))
        target = step.get("target", step.get("操作对象", ""))
        data = step.get("data", "")
        line = f"{seq}. {action}"
        if target:
            line += f" [{target}]"
        if data:
            line += f"（数据：{data}）"
        lines.append(line)
    return "\n".join(lines)


def _flatten_test_data(test_data: Any) -> str:
    """将测试数据转换为文本。"""
    if not test_data:
        return ""
    if isinstance(test_data, str):
        return test_data
    if isinstance(test_data, dict):
        lines = [f"{k}: {v}" for k, v in test_data.items()]
        return "\n".join(lines)
    if isinstance(test_data, list):
        lines = []
        for idx, item in enumerate(test_data, start=1):
            if isinstance(item, dict):
                lines.append("; ".join(f"{k}: {v}" for k, v in item.items()))
            else:
                lines.append(f"{idx}. {item}")
        return "\n".join(lines)
    return str(test_data)


def _flatten_list(items: list[str] | str | None) -> str:
    """将列表转换为数字序号文本。"""
    if not items:
        return ""
    if isinstance(items, str):
        return items
    lines = []
    for idx, result in enumerate(items, start=1):
        lines.append(f"{idx}. {result}")
    return "\n".join(lines)


def _extract_field(case: dict[str, Any], *keys: str, default: Any = "") -> Any:
    """从字典中按多个候选键提取第一个非空值。"""
    for key in keys:
        if key not in case:
            continue
        value = case[key]
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        if isinstance(value, (list, dict)) and not value:
            continue
        return value
    return default


def export_test_cases_to_docx(
    test_cases: list[dict[str, Any]],
    output_path: str | Path,
) -> str:
    """
    将测试用例列表导出为 Word (DOCX) 文件。

    支持字段：id, title, preconditions, steps, test_data, expected_results 等。

    Args:
        test_cases: 测试用例字典列表
        output_path: 导出的 DOCX 文件路径

    Returns:
        导出文件的绝对路径字符串。
    """
    if not test_cases:
        raise ValueError("测试用例列表为空，无法导出 DOCX。")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    
    # 列宽设定（4列基准）
    # 适当放宽第1列标签列，避免“测试用例名称/综述/通过准则”等字段在 Word 中被挤成竖向换行。
    total_width = 16.5
    col_widths = [
        total_width * 0.24,
        total_width * 0.28,
        total_width * 0.12,
        total_width * 0.36,
    ]
    
    for idx, case in enumerate(test_cases, start=1):
        case_id = _extract_field(case, "id", "case_id", "用例编号", "用例标识", default=f"TC-{idx:03d}")
        title = _extract_field(case, "title", "name", "case_name", "用例标题", "用例名称", "测试用例名称", default="")
        module = _extract_field(case, "module", "所属模块", "功能模块", "模块", "追踪关系", default="")
        remarks = _extract_field(case, "remarks", "备注", "关联需求", "用例说明", "测试用例综述", default="")
        preconditions = _flatten_list(_extract_field(case, "preconditions", "前置条件", "前提条件", "前提与约束", "前提和约束", default=None))
        design_method = _extract_field(case, "type", "case_type", "用例类型", "测试类型", "设计方法", default="")
        
        # 步骤需要特别处理（因为对应行）
        steps = _extract_field(case, "steps", "测试步骤", "操作步骤", "输入及操作", default=[])
        if isinstance(steps, str):
            steps = [steps]
        test_data = _flatten_test_data(_extract_field(case, "test_data", "testData", "测试数据", "具体测试数据", "测试输入", "输入数据", default=None))
        
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"用例 {idx}：{title}")
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = '宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        # 根据固定模板创建行
        num_fixed_rows_before = 8
        num_step_rows = max(len(steps), 1)
        num_fixed_rows_after = 7
        total_rows = num_fixed_rows_before + num_step_rows + num_fixed_rows_after
        
        table = doc.add_table(rows=total_rows, cols=4)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        tbl_pr = table._element.xpath('w:tblPr')
        if tbl_pr:
            tbl_w = parse_xml(f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>')
            tbl_layout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
            tbl_pr[0].append(tbl_w)
            tbl_pr[0].append(tbl_layout)
            
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)
                row.cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
        def set_row(row_idx, c0, c1, c2, c3, merge_1_to_3=False, merge_all=False, 
                    bold_0=True, bold_1=False, bold_2=False, bold_3=False):
            cells = table.rows[row_idx].cells
            
            def fill_cell(cell, text, is_bold):
                cell.text = str(text) if text else ""
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = '宋体'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                        run.font.size = Pt(10.5) # 5号字体
                        run.bold = is_bold
                        
            fill_cell(cells[0], c0, bold_0)
            
            if merge_all:
                cells[0].merge(cells[1]).merge(cells[2]).merge(cells[3])
                for p in cells[0].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif merge_1_to_3:
                fill_cell(cells[1], c1, bold_1)
                cells[1].merge(cells[2]).merge(cells[3])
            else:
                fill_cell(cells[1], c1, bold_1)
                fill_cell(cells[2], c2, bold_2)
                fill_cell(cells[3], c3, bold_3)
        
        # 填充前置固定部分
        set_row(0, "测试用例名称", title, "用例标识", case_id, bold_2=True)
        set_row(1, "追踪关系", module, "", "", merge_1_to_3=True)
        set_row(2, "测试用例综述", remarks, "", "", merge_1_to_3=True)
        set_row(3, "用例初始化", "", "", "", merge_1_to_3=True)
        set_row(4, "前提和约束", preconditions, "", "", merge_1_to_3=True)
        set_row(5, "设计方法", design_method, "", "", merge_1_to_3=True)
        set_row(6, "测试数据", test_data, "", "", merge_1_to_3=True)
        set_row(7, "序号", "输入及操作", "期望结果与评估标准", "实测结果", bold_1=True, bold_2=True, bold_3=True)
        
        # 居中显示步骤表头行
        for i in range(4):
            for p in table.rows[7].cells[i].paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        # 步骤动态行
        expected_list = _extract_field(case, "expected_results", "expected", "预期结果", "期望结果", "期望结果与评估标准", default=[])
        if isinstance(expected_list, str): expected_list = [expected_list]
        
        row_offset = 8
        if not steps:
            set_row(row_offset, "1", "", _flatten_list(expected_list), "", bold_0=False)
            for i in range(4):
                for p in table.rows[row_offset].cells[i].paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            for i, step in enumerate(steps):
                if isinstance(step, dict):
                    seq = step.get("seq", step.get("step", i + 1))
                    action = step.get("action", step.get("操作描述", ""))
                    target = step.get("target", step.get("操作对象", ""))
                    data = step.get("data", "")
                    exp = step.get("expected", step.get("预期结果", ""))
                else:
                    seq = i + 1
                    action = str(step)
                    target = ""
                    data = ""
                    exp = ""
                
                input_action = action
                if target: input_action += f" [{target}]"
                if data: input_action += f"\n(数据：{data})"
                
                if not exp:
                    if len(expected_list) == len(steps):
                        # 数量一致时，完美1对1映射
                        exp = expected_list[i]
                    else:
                        # 数量不匹配时，大模型通常是把断言全堆在了数组最后。
                        # 因此：前置步骤采用基础规则判定保底，所有断言垒到最后一步。
                        if i == len(steps) - 1:
                            exp = _flatten_list(expected_list)
                            if not exp: exp = "操作成功，预期符合系统设定行为"
                        else:
                            # 启发式规则保底：根据用户的输入动词生成合理的默认预期验证
                            action_text = str(action).lower()
                            if "输入" in action_text or "填写" in action_text:
                                exp = "成功输入并显示对应内容"
                            elif "选择" in action_text or "下拉" in action_text or "字典" in action_text:
                                exp = "选择字典成功，正常显示所选项"
                            elif "打开" in action_text or "进入" in action_text or "跳转" in action_text:
                                exp = "页面加载成功，正常展示界面元素"
                            elif "上传" in action_text or "导入" in action_text:
                                exp = "文件选择成功，触发上传交互"
                            elif "勾选" in action_text or "点击" in action_text:
                                exp = "交互响应正常"
                            else:
                                exp = "操作执行成功"
                            
                set_row(row_offset + i, seq, input_action, exp, "", bold_0=False)
                for p in table.rows[row_offset + i].cells[0].paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        # 填充后置固定部分
        after_offset = row_offset + num_step_rows
        set_row(after_offset    , "测试用例终止条件", "正常终止条件为功能实现，异常终止条件为功能未实现", "", "", merge_1_to_3=True)
        set_row(after_offset + 1, "测试用例通过准则", "与期望结果一致", "", "", merge_1_to_3=True)
        set_row(after_offset + 2, "执行状态", "", "执行结果", "", bold_2=True)
        set_row(after_offset + 3, "问题报告单标识", "", "", "", merge_1_to_3=True)
        set_row(after_offset + 4, "测试人员", "", "", "", merge_1_to_3=True)
        set_row(after_offset + 5, "测试时间", "", "", "", merge_1_to_3=True)
        set_row(after_offset + 6, "未执行原因", "", "", "", merge_1_to_3=True)
        
        # 全部居中和靠左微调
        # 这里对整个表格做一次默认遍历：表头和左侧标签列居中，输入框靠左
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    if r_idx in [6, 7] or c_idx == 0 or (r_idx in [0, after_offset+2] and c_idx == 2):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        doc.add_paragraph()
        
    try:
        doc.save(str(output_path))
    except PermissionError:
        # 如果文件被占用（如用户正在Word中打开），自动加时间戳换一个名字重试
        import time
        new_name = f"{output_path.stem}_{int(time.time())}{output_path.suffix}"
        output_path = output_path.with_name(new_name)
        doc.save(str(output_path))
        
    return str(output_path.resolve())
