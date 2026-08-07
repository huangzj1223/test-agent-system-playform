
"""Project-level export helpers for API and scenario test artifacts."""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import datetime
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config.minio_client import MinIOClient, MinIOError
from app.models.api_endpoint import APIEndpoint
from app.models.api_test import APITest
from app.models.attachment import Attachment, AttachmentEntityType
from app.models.project import Project
from app.models.test_scenario import TestScenario


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ZIP_CONTENT_TYPE = "application/zip"


def _safe_filename(value: str | None, fallback: str = "untitled") -> str:
    raw = (value or fallback).strip() or fallback
    raw = re.sub(r"[\\/:*?\"<>|\r\n\t]+", "_", raw)
    raw = re.sub(r"\s+", " ", raw).strip(" .")
    return raw[:120] or fallback


def _json_dumps(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2)


def _normalize_priority(value: Any, default: str = "P1") -> str:
    if value is None or value == "":
        return default
    raw = str(value).strip()
    compact = raw.replace(" ", "").upper()
    if compact in {"P0", "P1", "P2", "P3", "P4"}:
        return compact
    mapping = {
        "CRITICAL": "P0",
        "BLOCKER": "P0",
        "URGENT": "P0",
        "HIGHEST": "P0",
        "\u7d27\u6025": "P0",
        "\u4e25\u91cd": "P0",
        "\u6700\u9ad8": "P0",
        "HIGH": "P1",
        "\u9ad8": "P1",
        "\u91cd\u8981": "P1",
        "MEDIUM": "P2",
        "MID": "P2",
        "NORMAL": "P2",
        "\u4e2d": "P2",
        "\u666e\u901a": "P2",
        "LOW": "P3",
        "\u4f4e": "P3",
        "\u6700\u4f4e": "P3",
    }
    return mapping.get(compact, mapping.get(raw, default))


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize_api_cases(payload: Any, endpoint: APIEndpoint | None = None) -> list[dict[str, Any]]:
    if isinstance(payload, dict):
        raw_cases = payload.get("test_cases") or payload.get("cases") or payload.get("data") or []
    elif isinstance(payload, list):
        raw_cases = payload
    else:
        raw_cases = []

    cases: list[dict[str, Any]] = []
    for index, item in enumerate(raw_cases, start=1):
        if not isinstance(item, dict):
            item = {"title": str(item)}

        steps_value = item.get("steps") or item.get("test_steps") or item.get("actions") or []
        expected_value = item.get("expected_result") or item.get("expected") or item.get("expected_results") or item.get("assertions") or ""

        if isinstance(steps_value, list):
            steps = []
            for step_index, step in enumerate(steps_value, start=1):
                if isinstance(step, dict):
                    steps.append(step.get("action") or step.get("step") or step.get("description") or _json_dumps(step))
                else:
                    steps.append(str(step))
            steps_text = "\n".join(f"{i}. {step}" for i, step in enumerate(steps, start=1))
        else:
            steps_text = str(steps_value or "")

        if isinstance(expected_value, list):
            expected_text = "\n".join(str(v) if not isinstance(v, dict) else _json_dumps(v) for v in expected_value)
        else:
            expected_text = _json_dumps(expected_value)

        endpoint_name = endpoint.display_name if endpoint else item.get("endpoint") or item.get("module") or ""
        cases.append({
            "case_id": item.get("id") or item.get("case_id") or item.get("identifier") or f"API-TC-{index:03d}",
            "module": endpoint_name,
            "title": item.get("title") or item.get("name") or item.get("case_name") or f"Test Case {index}",
            "priority": _normalize_priority(item.get("priority")),
            "type": item.get("type") or item.get("case_type") or "API",
            "preconditions": item.get("preconditions") or item.get("precondition") or "",
            "steps": steps_text,
            "expected_result": expected_text,
            "source": endpoint_name,
        })
    return cases


def _workbook_to_bytes(rows: list[dict[str, Any]], headers: list[tuple[str, str]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for column_index, (_, label) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column_index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, _) in enumerate(headers, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=row.get(key, ""))
            cell.alignment = wrap

    for column_index, (key, label) in enumerate(headers, start=1):
        max_length = len(label)
        for row in rows:
            max_length = max(max_length, min(len(str(row.get(key, "") or "")), 80))
        ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 60)

    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class TestArtifactExportService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def _get_project(self, project_identifier: str) -> Project:
        project_id: UUID | None = None
        try:
            project_id = UUID(project_identifier)
        except ValueError:
            project_id = None

        stmt = select(Project)
        if project_id:
            stmt = stmt.where(Project.id == project_id)
        else:
            stmt = stmt.where(Project.identifier == project_identifier)
        project = (await self.db.execute(stmt)).scalar_one_or_none()
        if not project:
            raise ValueError(f"Project {project_identifier} not found")
        return project

    async def export_api_test_cases_excel(self, project_identifier: str) -> tuple[bytes, str, str]:
        project = await self._get_project(project_identifier)
        stmt = (
            select(Attachment, APIEndpoint)
            .join(APIEndpoint, APIEndpoint.id == Attachment.entity_id, isouter=True)
            .where(Attachment.project_id == project.id)
            .where(Attachment.entity_type == AttachmentEntityType.API_TEST_CASE)
            .order_by(Attachment.created_at.asc())
        )
        result = await self.db.execute(stmt)

        rows: list[dict[str, Any]] = []
        for attachment, endpoint in result.all():
            try:
                payload = json.loads(_decode_bytes(MinIOClient.download_file(attachment.object_name)))
            except Exception as exc:
                rows.append({
                    "case_id": str(attachment.id),
                    "module": endpoint.display_name if endpoint else attachment.file_name,
                    "title": "Failed to parse test case file",
                    "priority": "",
                    "type": "API",
                    "preconditions": "",
                    "steps": "",
                    "expected_result": str(exc),
                    "source": attachment.object_name,
                })
                continue
            rows.extend(_normalize_api_cases(payload, endpoint))

        headers = [
            ("case_id", "\u7528\u4f8b\u7f16\u53f7"),
            ("module", "\u63a5\u53e3/\u6a21\u5757"),
            ("title", "\u7528\u4f8b\u6807\u9898"),
            ("priority", "\u4f18\u5148\u7ea7"),
            ("type", "\u7528\u4f8b\u7c7b\u578b"),
            ("preconditions", "\u524d\u7f6e\u6761\u4ef6"),
            ("steps", "\u6d4b\u8bd5\u6b65\u9aa4"),
            ("expected_result", "\u9884\u671f\u7ed3\u679c"),
            ("source", "\u6765\u6e90"),
        ]
        content = _workbook_to_bytes(rows, headers, "Test Cases")
        filename = f"api-test-cases-{project.identifier}-{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        return content, filename, EXCEL_CONTENT_TYPE

    async def export_api_test_scripts_zip(self, project_identifier: str) -> tuple[bytes, str, str]:
        project = await self._get_project(project_identifier)
        zip_buffer = io.BytesIO()
        seen_paths: set[str] = set()
        count = 0

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            tests = (await self.db.execute(
                select(APITest).where(APITest.project_id == project.id).order_by(APITest.created_at.asc())
            )).scalars().all()
            for test in tests:
                if not test.script_path or test.script_path in seen_paths:
                    continue
                seen_paths.add(test.script_path)
                try:
                    data = MinIOClient.download_file(test.script_path)
                except Exception:
                    continue
                extension = _script_extension(test.script_language, test.script_path)
                name = _safe_filename(f"{test.identifier}-{test.name}")
                zf.writestr(f"api-tests/{name}.{extension}", data)
                count += 1

            attachment_stmt = (
                select(Attachment, APIEndpoint)
                .join(APIEndpoint, APIEndpoint.id == Attachment.entity_id, isouter=True)
                .where(Attachment.project_id == project.id)
                .where(Attachment.entity_type == AttachmentEntityType.API_TEST_SCRIPT)
                .order_by(Attachment.created_at.asc())
            )
            for attachment, endpoint in (await self.db.execute(attachment_stmt)).all():
                if attachment.object_name in seen_paths:
                    continue
                seen_paths.add(attachment.object_name)
                try:
                    data = MinIOClient.download_file(attachment.object_name)
                except Exception:
                    continue
                suffix = PathSuffix.from_name(attachment.file_name)
                name = _safe_filename(endpoint.display_name if endpoint else attachment.file_name)
                zf.writestr(f"api-test-artifacts/{name}.{suffix}", data)
                count += 1

            if count == 0:
                zf.writestr("README.txt", "No test scripts are available for export.")

        filename = f"api-test-scripts-{project.identifier}-{datetime.now().strftime('%Y%m%d%H%M%S')}.zip"
        return zip_buffer.getvalue(), filename, ZIP_CONTENT_TYPE

    async def export_scenario_test_cases_excel(self, project_identifier: str) -> tuple[bytes, str, str]:
        project = await self._get_project(project_identifier)
        scenarios = (await self.db.execute(
            select(TestScenario)
            .options(selectinload(TestScenario.steps))
            .where(TestScenario.project_id == project.id)
            .order_by(TestScenario.created_at.asc())
        )).scalars().all()

        rows: list[dict[str, Any]] = []
        for scenario in scenarios:
            steps = sorted(scenario.steps or [], key=lambda step: step.step_order)
            rows.append({
                "case_id": scenario.identifier,
                "module": "Scenario Test",
                "title": scenario.name,
                "priority": _normalize_priority("P1"),
                "type": "SCENARIO",
                "preconditions": _json_dumps(scenario.global_variables),
                "steps": "\n".join(f"{step.step_order}. {step.name}" for step in steps),
                "expected_result": "\n".join(_format_step_assertions(step.assertions) for step in steps if step.assertions),
                "source": str(scenario.id),
            })

        headers = [
            ("case_id", "\u7528\u4f8b\u7f16\u53f7"),
            ("module", "\u6a21\u5757"),
            ("title", "\u573a\u666f/\u7528\u4f8b\u6807\u9898"),
            ("priority", "\u4f18\u5148\u7ea7"),
            ("type", "\u7528\u4f8b\u7c7b\u578b"),
            ("preconditions", "\u524d\u7f6e\u6761\u4ef6/\u53d8\u91cf"),
            ("steps", "\u6d4b\u8bd5\u6b65\u9aa4"),
            ("expected_result", "\u9884\u671f\u7ed3\u679c/\u65ad\u8a00"),
            ("source", "\u573a\u666fID"),
        ]
        content = _workbook_to_bytes(rows, headers, "Test Cases")
        filename = f"scenario-test-cases-{project.identifier}-{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        return content, filename, EXCEL_CONTENT_TYPE

    async def export_scenario_test_scripts_zip(self, project_identifier: str) -> tuple[bytes, str, str]:
        project = await self._get_project(project_identifier)
        scenarios = (await self.db.execute(
            select(TestScenario)
            .options(selectinload(TestScenario.steps))
            .where(TestScenario.project_id == project.id)
            .order_by(TestScenario.created_at.asc())
        )).scalars().all()

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            if not scenarios:
                zf.writestr("README.txt", "No test scripts are available for export.")
            for scenario in scenarios:
                script = _build_scenario_playwright_script(scenario)
                name = _safe_filename(f"{scenario.identifier}-{scenario.name}")
                zf.writestr(f"scenario-tests/{name}.spec.ts", script.encode("utf-8"))

        filename = f"scenario-test-scripts-{project.identifier}-{datetime.now().strftime('%Y%m%d%H%M%S')}.zip"
        return zip_buffer.getvalue(), filename, ZIP_CONTENT_TYPE


class PathSuffix:
    @staticmethod
    def from_name(name: str | None) -> str:
        safe = name or "script.txt"
        if "." in safe:
            suffix = safe.rsplit(".", 1)[-1].lower()
            if re.fullmatch(r"[a-z0-9]{1,8}", suffix):
                return suffix
        return "txt"


def _script_extension(language: str | None, path: str | None) -> str:
    if path and "." in path.rsplit("/", 1)[-1]:
        suffix = path.rsplit(".", 1)[-1].lower()
        if re.fullmatch(r"[a-z0-9]{1,8}", suffix):
            return suffix
    return {
        "typescript": "ts",
        "javascript": "js",
        "python": "py",
        "java": "java",
    }.get((language or "").lower(), "txt")


def _format_step_assertions(assertions: Any) -> str:
    if not assertions:
        return ""
    if isinstance(assertions, list):
        return "; ".join(_json_dumps(item) for item in assertions)
    return _json_dumps(assertions)


def _build_scenario_playwright_script(scenario: TestScenario) -> str:
    steps = sorted(scenario.steps or [], key=lambda step: step.step_order)
    lines = [
        "import { test, expect, request } from '@playwright/test';",
        "",
        f"test({json.dumps(scenario.name, ensure_ascii=False)}, async ({{ request }}) => {{",
        "  const baseURL = process.env.API_BASE_URL || '';",
        "  const variables: Record<string, unknown> = {};",
    ]

    if scenario.global_variables:
        lines.append(f"  Object.assign(variables, {json.dumps(scenario.global_variables, ensure_ascii=False, indent=2)});")

    for step in steps:
        request_override = step.request_override or {}
        headers_override = step.headers_override or {}
        method = str(request_override.get("method") or "GET").lower()
        path = request_override.get("path") or request_override.get("url") or ""
        body = request_override.get("body") or request_override.get("json")
        lines.extend([
            "",
            f"  // Step {step.step_order}: {step.name}",
            f"  const response{step.step_order} = await request.{method}(baseURL + {json.dumps(path, ensure_ascii=False)}, {{",
        ])
        if headers_override:
            lines.append(f"    headers: {json.dumps(headers_override, ensure_ascii=False, indent=4)},")
        if body is not None:
            lines.append(f"    data: {json.dumps(body, ensure_ascii=False, indent=4)},")
        lines.append("  });")
        assertions = step.assertions or []
        status_assertion = next((item for item in assertions if isinstance(item, dict) and item.get("type") == "status"), None)
        if status_assertion and status_assertion.get("expected") is not None:
            lines.append(f"  expect(response{step.step_order}.status()).toBe({json.dumps(status_assertion.get('expected'))});")
        else:
            lines.append(f"  expect(response{step.step_order}.ok()).toBeTruthy();")
        if step.extractors:
            lines.append(f"  // Extractors: {json.dumps(step.extractors, ensure_ascii=False)}")

    lines.append("});")
    lines.append("")
    return "\n".join(lines)
